#!/usr/bin/env python3
"""Generate the report, embedding only an actual screenshot supplied by runtime capture."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "submission"
SCREENSHOT = SUBMISSION / "screenshots" / "location-tracker-running.png"
OUTPUT = SUBMISSION / "MSCS533_Final_Project_LocationTracker_AM.docx"
REPOSITORY = "https://github.com/AshishM26/MSCS-533-Project_2_AM"


def add_link(paragraph, label: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "512BD4")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_code(document, source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(source.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def main() -> None:
    status = json.loads((SUBMISSION / "validation.json").read_text())
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor.from_string("512BD4")
    document.core_properties.title = 'Multiplatform "Location Tracker" App with .NET MAUI'
    document.core_properties.author = "Ashish Bhanudas Mahajan"
    document.core_properties.subject = "MSCS-533-A01 Final Software Development Project"

    document.add_paragraph("UNIVERSITY OF THE CUMBERLANDS", "Subtitle")
    document.add_heading("Final Software Development Project", 0)
    document.add_paragraph('Multiplatform "Location Tracker" App with .NET MAUI', "Subtitle")
    for line in (
        "Ashish Bhanudas Mahajan",
        "MSCS-533-A01",
        "Software Engineering and Multiplatform App Development",
        "Fall 2026 First Bi-Term",
        f"Validation date: {status['date']}",
    ):
        document.add_paragraph(line)
    if not SCREENSHOT.exists() or not status["runtime_complete"]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("DRAFT — final map runtime evidence and application screenshot are pending.")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("9E2040")
    add_link(document.add_paragraph(), REPOSITORY, REPOSITORY)
    document.add_page_break()

    document.add_heading("1. Project Overview", level=1)
    document.add_paragraph(
        "Location Tracker is a C# and .NET MAUI application that records geographic location samples while "
        "the app is in the foreground, stores them in SQLite, and extends the native Maps component with "
        "a heat-map-style density layer. The design separates location acquisition, persistence, aggregation, "
        "and presentation so each responsibility can be understood and maintained independently."
    )

    document.add_heading("2. Application Requirements Implemented", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Requirement"
    table.rows[0].cells[1].text = "Implementation"
    for requirement, implementation in (
        ("Location tracking", "Start/Stop commands; sequential GPS requests; cancellation on stop/background."),
        ("SQLite history", "Asynchronous local saves, ordered reload, count, and confirmed deletion."),
        ("Map and heat layer", "MAUI Maps Circle elements grouped by geographic density; latest-sample pin."),
        ("Usable controls", "Status, sample count, latest reading, centering, and disabled invalid actions."),
        ("Permissions and failures", "While-in-use permission; GPS/permission/storage messages; no background access."),
        ("Multiplatform source", "Android and iOS projects; Mac Catalyst retained on macOS."),
    ):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = requirement, implementation
    document.add_paragraph("Implementation does not by itself establish runtime acceptance. Verified results and remaining limitations appear in Section 8.")

    document.add_heading("3. Technologies Used", level=1)
    document.add_paragraph(
        ".NET SDK 10.0.302; C# with nullable reference types; .NET MAUI Controls and Maps 10.0.100; "
        "sqlite-net-pcl 1.11.285; xUnit; built-in dependency injection; and VS Code with Microsoft's MAUI, "
        "C# Dev Kit, and C# extensions. No .NET 11 preview is targeted."
    )

    document.add_heading("4. Application Architecture", level=1)
    document.add_paragraph(
        "The solution has three projects: LocationTracker provides the native UI and platform adapter; "
        "LocationTracker.Core contains portable models, data access, sampling orchestration, and heat-map logic; "
        "LocationTracker.Tests verifies behavior without a phone. MainViewModel exposes observable state and "
        "commands. The page handles dialogs and native map elements. MauiProgram.cs registers services and calls UseMauiMaps()."
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Source location"
    table.rows[0].cells[1].text = "Responsibility"
    for path, purpose in (
        ("LocationTracker/MauiProgram.cs", "Composition, SQLite path, and Maps setup"),
        ("LocationTracker/Views/MainPage.xaml(.cs)", "Layout, map circles, centering, and dialog"),
        ("LocationTracker/ViewModels/MainViewModel.cs", "State, commands, history and lifecycle"),
        ("LocationTracker/Services/MauiLocationProvider.cs", "Native GPS and permission requests"),
        ("LocationTracker.Core/Services/LocationTrackingService.cs", "Cancellable sampling and saving"),
        ("LocationTracker.Core/Data/LocationDatabase.cs", "Asynchronous SQLite operations"),
        ("LocationTracker.Core/Services/HeatMapService.cs", "Geographic grouping and density"),
        ("LocationTracker/Platforms/Android/AndroidManifest.xml", "Android location and map metadata"),
        ("LocationTracker/Platforms/iOS/Info.plist", "iOS location purpose description"),
    ):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = path, purpose

    document.add_heading("5. Location Tracking and SQLite Persistence", level=1)
    document.add_paragraph(
        "The service requests while-in-use permission, then asks MAUI Geolocation for a high-accuracy reading "
        "with a 15-second timeout. Each completed attempt is followed by an eight-second delay, preventing "
        "overlapping GPS requests. A CancellationToken stops pending requests and delays. Stop waits for an "
        "in-flight save to finish. Backgrounding the app cancels tracking and does not automatically restart it."
    )
    document.add_paragraph(
        "LocationRecord stores Id, Latitude, Longitude, TimestampUtc, and AccuracyMeters. The database is created "
        "lazily under FileSystem.AppDataDirectory as locations.db3. Asynchronous operations are serialized. "
        "Startup restores ordered history; Clear History requires confirmation and removes records, circles, and the latest pin. "
        "The database is local and unencrypted. The app does not upload history; operating-system backups and native map network services are separate considerations."
    )
    document.add_paragraph("C# source excerpt — LocationTracker.Core/Data/LocationDatabase.cs:")
    database_source = (ROOT / "LocationTracker.Core/Data/LocationDatabase.cs").read_text()
    start = database_source.index("    public async Task SaveLocationAsync")
    end = database_source.index("    public async Task<IReadOnlyList", start)
    add_code(document, database_source[start:end])

    document.add_heading("6. Heat Map Implementation", level=1)
    document.add_paragraph(
        "HeatMapService groups valid points into roughly 75-meter cells, adjusting longitude divisions by latitude. "
        "Each cell uses the mean latitude and circular mean longitude of its samples. Density equals the cell's count "
        "divided by the largest cell count. The page renders blue, amber, or red circles with increasing opacity and radius. "
        "The grid is coarsened if needed to keep at most 400 map elements, preserving all sample counts."
    )
    document.add_paragraph(
        "The overlay is a discrete relative-density visualization rather than a continuous kernel-density raster. "
        "Equal counts produce equal colors; longer dwell time produces denser cells. Grid boundaries may separate "
        "nearby points. The classroom implementation loads all saved history in memory, so unlimited long-term recording would need paging or retention controls."
    )

    document.add_heading("7. Platform Permissions / Manifest Configuration", level=1)
    document.add_paragraph(
        "AndroidManifest.xml declares ACCESS_COARSE_LOCATION and ACCESS_FINE_LOCATION, plus internet/network access. "
        "It contains Google Maps metadata with a build-time placeholder. A real key is supplied only through the ignored "
        "GoogleMaps.local.props file. Android backups are disabled. No background-location permission is requested."
    )
    document.add_paragraph(
        "Platforms/iOS/Info.plist defines NSLocationWhenInUseUsageDescription: “Location is used to record your route "
        "and visualize your location history on the map while the app is open.” Mac Catalyst includes the corresponding "
        "usage description and sandbox location entitlement. iOS Maps does not require the Android Google Maps key."
    )
    document.add_paragraph("Android manifest source — no real key is included:")
    add_code(document, (ROOT / "LocationTracker/Platforms/Android/AndroidManifest.xml").read_text())

    document.add_heading("8. Application Output", level=1)
    document.add_paragraph(status["tests"])
    document.add_paragraph(status["builds"])
    document.add_paragraph(status["runtime"])
    if SCREENSHOT.exists():
        document.add_picture(str(SCREENSHOT), width=Inches(3.1))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(status.get("screenshot_caption", "Figure 1. Actual application output captured from the running simulator."), "Caption")
    else:
        document.add_paragraph(
            "Screenshot pending: no final map-output screenshot has been captured. Install and configure Xcode "
            "and an iOS simulator, execute the acceptance checks in submission/RUNTIME_GUIDE.md, and capture "
            "submission/screenshots/location-tracker-running.png. Rerun scripts/create_submission.py to embed it. "
            "This draft must not be submitted as proof of completed runtime validation."
        )

    document.add_heading("9. GitHub Repository", level=1)
    add_link(document.add_paragraph(), REPOSITORY, REPOSITORY)
    document.add_paragraph(
        "The repository provides the complete C# source, MAUI project, platform manifests, automated tests, "
        "README, and validation instructions. Build output, simulator data, local keys, and private machine paths "
        "are excluded from version control. Source code is shared through GitHub rather than reproduced in full here."
    )

    document.add_heading("10. Conclusion", level=1)
    document.add_paragraph(
        "The implementation brings together foreground geolocation, asynchronous SQLite persistence, and a native "
        "map density overlay using a maintainable MVVM design. Automated tests verify portable behavior, while "
        "device-specific permission, map rendering, and lifecycle checks are tracked separately. "
        + ("The recorded runtime checks and screenshot provide the application-output evidence." if status["runtime_complete"] and SCREENSHOT.exists()
           else "Final submission readiness remains conditional on completing the map runtime checks and inserting the real screenshot.")
    )

    document.add_heading("References", level=1)
    for title, url in (
        ("Microsoft Learn: MAUI Maps", "https://learn.microsoft.com/en-us/dotnet/maui/user-interface/controls/map?view=net-maui-10.0"),
        ("Microsoft Learn: MAUI geolocation", "https://learn.microsoft.com/en-us/dotnet/maui/platform-integration/device/geolocation?view=net-maui-10.0"),
        ("Microsoft Learn: C# coding conventions", "https://learn.microsoft.com/en-us/dotnet/csharp/fundamentals/coding-style/coding-conventions"),
    ):
        add_link(document.add_paragraph(), title, url)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("MSCS-533-A01  |  Ashish Bhanudas Mahajan  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    expected_images = 1 if SCREENSHOT.exists() else 0
    assert len(reopened.inline_shapes) == expected_images
    assert all(any(p.text.startswith(f"{index}.") for p in reopened.paragraphs) for index in range(1, 11))
    print(f"Created {OUTPUT.relative_to(ROOT)}; verified 10 sections and {expected_images} embedded screenshot(s).")


if __name__ == "__main__":
    main()
